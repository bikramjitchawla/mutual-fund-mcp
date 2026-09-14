"""Public PPFAS monthly XLSX disclosures, mapped to verified Direct Growth codes."""

from __future__ import annotations

import calendar
import io
import math
import re
import time
from datetime import date, datetime
from html.parser import HTMLParser
from threading import RLock
from urllib.parse import urljoin, urlsplit
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from clients.public_http import fetch_bytes
from services.errors import FundError

BASE = "https://amc.ppfas.com"
INDEX = BASE + "/downloads/portfolio-disclosure/index.php"
# Portfolios are scheme-level (shared by plans), represented once to avoid double counting.
FUNDS = {
    "122639": ("PPFCF", "Parag Parikh Flexi Cap Fund"),
    "143269": ("PPLF", "Parag Parikh Liquid Fund"),
    "147481": ("PPTSF", "Parag Parikh ELSS Tax Saver Fund"),
    "148958": ("PPCHF", "Parag Parikh Conservative Hybrid Fund"),
}
SCOPE = "Four PPFAS schemes only, represented by Direct Growth codes; ISIN-bearing cash securities only. Cash, receivables and derivatives excluded. Includes arbitrage cash legs; weights are not net directional exposure. Not an all-AMC search."


def validate_isin(value: str) -> str:
    value = value.strip().upper()
    if not re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", value):
        raise FundError(
            "INVALID_ISIN", "ISIN must be a valid 12-character security identifier."
        )
    digits = "".join(str(ord(c) - 55) if c.isalpha() else c for c in value)
    total = 0
    for i, char in enumerate(reversed(digits)):
        number = int(char) * (2 if i % 2 else 1)
        total += number // 10 + number % 10
    if total % 10:
        raise FundError("INVALID_ISIN", "ISIN checksum is invalid.")
    return value


def month_end(month: str) -> date:
    try:
        if not re.fullmatch(r"[0-9]{4}-[0-9]{2}", month):
            raise ValueError
        first = date.fromisoformat(month + "-01")
        return first.replace(day=calendar.monthrange(first.year, first.month)[1])
    except (TypeError, ValueError) as error:
        raise FundError("INVALID_MONTH", "Month must use YYYY-MM.") from error


class _Links(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            self.links.extend(value for key, value in attrs if key == "href" and value)


class PPFASProvider:
    def __init__(self, fetcher=fetch_bytes):
        self.fetcher = fetcher
        self._links = None
        self._expires = 0.0
        self._lock = RLock()

    def disclosures(self) -> dict[tuple[str, str], str]:
        with self._lock:
            return self._disclosures()

    def _disclosures(self) -> dict[tuple[str, str], str]:
        if self._links is not None and time.monotonic() < self._expires:
            return dict(self._links)
        parser = _Links()
        try:
            parser.feed(self.fetcher(INDEX).decode("utf-8"))
        except UnicodeError as error:
            raise FundError(
                "INVALID_PROVIDER_RESPONSE", "Invalid disclosure index encoding."
            ) from error
        links = {}
        prefixes = {value[0]: code for code, value in FUNDS.items()}
        for href in parser.links:
            url = urljoin(BASE, href)
            parts = urlsplit(url)
            if parts.scheme != "https" or parts.netloc != "amc.ppfas.com":
                continue
            match = re.fullmatch(
                r"/downloads/portfolio-disclosure/([0-9]{4})/([A-Z]+)_PPFAS_Monthly_Portfolio_Report_([A-Za-z]+)_([0-9]{1,2})_([0-9]{4})\.xlsx",
                parts.path,
            )
            if not match or match[2] not in prefixes or match[1] != match[5]:
                continue
            try:
                day = datetime.strptime(
                    f"{match[3]} {match[4]} {match[5]}", "%B %d %Y"
                ).date()
            except ValueError:
                continue
            if day != month_end(day.strftime("%Y-%m")):
                continue
            links[(prefixes[match[2]], day.strftime("%Y-%m"))] = url
        if not links:
            raise FundError(
                "INVALID_PROVIDER_RESPONSE",
                "No supported PPFAS XLSX disclosure links found.",
            )
        self._links, self._expires = links, time.monotonic() + 3600
        return dict(links)

    def get_snapshot(self, scheme_code: str, month: str, url: str) -> dict:
        if self.disclosures().get((scheme_code, month)) != url:
            raise FundError(
                "NO_HOLDINGS_DATA", "No official disclosure for this fund and month."
            )
        return parse_workbook(self.fetcher(url), scheme_code, month, url)


def parse_workbook(content: bytes, scheme_code: str, month: str, source: str) -> dict:
    if scheme_code not in FUNDS:
        raise FundError(
            "UNSUPPORTED_HOLDINGS_SCHEME", "Scheme is outside PPFAS holdings coverage."
        )
    expected_date = month_end(month)
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds size limit")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        try:
            prefix, name = FUNDS[scheme_code]
            if prefix not in workbook.sheetnames:
                raise ValueError("Missing expected scheme sheet")
            sheet = workbook[prefix]
            if sheet.max_row > 10000 or sheet.max_column > 100:
                raise ValueError("Unexpected workbook dimensions")
            rows = list(sheet.iter_rows())
            header_text = " ".join(str(c.value or "") for row in rows[:4] for c in row)
            # Former ELSS name appears in older disclosures.
            names = (
                [name, "Parag Parikh Tax Saver Fund"] if prefix == "PPTSF" else [name]
            )
            if not any(n in header_text for n in names):
                raise ValueError("Mismatched fund title")
            date_match = re.search(
                r"as on\s+([A-Za-z]+ \d{1,2},? \d{4})", header_text, re.I
            )
            if (
                not date_match
                or datetime.strptime(date_match[1].replace(",", ""), "%B %d %Y").date()
                != expected_date
            ):
                raise ValueError("Mismatched portfolio date")
            header_index = next(
                i
                for i, row in enumerate(rows[:15])
                if len(row) >= 7 and row[2].value == "ISIN"
            )
            headers = [
                " ".join(str(c.value or "").split()).lower() for c in rows[header_index]
            ]
            if (
                headers[1] != "name of the instrument"
                or headers[4] != "quantity"
                or "lakhs" not in headers[5]
                or "net" not in headers[6]
            ):
                raise ValueError("Unsupported disclosure columns")
            holdings = {}
            found_total = False
            for row in rows[header_index + 1 :]:
                label = str(row[1].value or "").strip()
                if label.upper() == "GRAND TOTAL":
                    found_total = True
                    break
                raw_isin = row[2].value
                if not raw_isin:
                    continue
                isin = validate_isin(str(raw_isin))
                quantity = float(row[4].value)
                market_value = float(row[5].value) * 100000
                cell = row[6]
                if isinstance(cell.value, str):
                    weight = float(cell.value.replace("$", "").replace("%", "").strip())
                else:
                    weight = float(cell.value) * (
                        100 if "%" in cell.number_format else 1
                    )
                if (
                    not label
                    or any(
                        not math.isfinite(v) or v < 0
                        for v in (quantity, market_value, weight)
                    )
                    or weight > 100
                ):
                    raise ValueError("Invalid holding values")
                if isin in holdings:
                    for key, value in (
                        ("quantity", quantity),
                        ("market_value", market_value),
                        ("portfolio_weight_pct", weight),
                    ):
                        holdings[isin][key] += value
                else:
                    holdings[isin] = {
                        "isin": isin,
                        "security_name": label,
                        "sector": str(row[3].value or "").strip() or None,
                        "quantity": quantity,
                        "market_value": market_value,
                        "portfolio_weight_pct": weight,
                    }
            if not found_total or not holdings:
                raise ValueError("Incomplete or empty disclosure")
            for item in holdings.values():
                item["portfolio_weight_pct"] = round(item["portfolio_weight_pct"], 6)
            return {
                "scheme_code": scheme_code,
                "scheme_name": name + " - Direct Plan - Growth",
                "portfolio_date": expected_date.isoformat(),
                "source": source,
                "holdings": list(holdings.values()),
                "scope": SCOPE,
            }
        finally:
            workbook.close()
    except (
        ValueError,
        TypeError,
        KeyError,
        IndexError,
        StopIteration,
        BadZipFile,
        FundError,
        ParseError,
        OverflowError,
    ) as error:
        raise FundError(
            "INVALID_PROVIDER_RESPONSE",
            "PPFAS workbook failed schema, identity or value validation.",
        ) from error
